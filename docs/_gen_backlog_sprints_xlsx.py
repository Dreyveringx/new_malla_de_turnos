# -*- coding: utf-8 -*-
"""Genera BACKLOG-SPRINTS Excel (hojas Sprint 1..9) + regenera el .md alineado."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DOCS = Path(__file__).resolve().parent
XLSX = DOCS / "BACKLOG-SPRINTS-MALLA-TURNOS.xlsx"
MD = DOCS / "BACKLOG-SPRINTS-MALLA-TURNOS.md"

# Capacidad fija por rol
CAP = 80


def activities(sprint_n: int, hu_range: str, focus: str) -> list[tuple[str, int, str]]:
    """Fullstack diseña BD solo en Sprint 1. Backend no hace Flyway."""
    if sprint_n == 1:
        fs = [
            (
                "Diseño de la base de datos del módulo (modelo lógico, tablas iniciales, índices, "
                "constraints tenant) + primeras migraciones Flyway",
                16,
                "Fullstack",
            ),
            (
                f"Arquitectura/diseño técnico del alcance {hu_range} y contratos de API",
                10,
                "Fullstack",
            ),
            (
                f"Implementación núcleo backend de {focus} ({hu_range})",
                18,
                "Fullstack",
            ),
            (
                f"Pieza FE compleja o patrón reutilizable del sprint ({hu_range})",
                14,
                "Fullstack",
            ),
            ("Revisión PR aprendices + pair BE/FE", 10, "Fullstack"),
            ("Pruebas de integración y ajuste de criterios con QA", 8, "Fullstack"),
            ("Hardening / deuda técnica del sprint", 4, "Fullstack"),
        ]
    else:
        fs = [
            (
                f"Arquitectura/diseño técnico del alcance {hu_range}, contratos de API "
                f"y migraciones Flyway del sprint (si el esquema crece)",
                12,
                "Fullstack",
            ),
            (
                f"Implementación núcleo backend de {focus} ({hu_range})",
                22,
                "Fullstack",
            ),
            (
                f"Pieza FE compleja o patrón reutilizable del sprint ({hu_range})",
                16,
                "Fullstack",
            ),
            ("Revisión PR aprendices + pair BE/FE", 14, "Fullstack"),
            ("Pruebas de integración y ajuste de criterios con QA", 10, "Fullstack"),
            ("Hardening / deuda técnica del sprint", 6, "Fullstack"),
        ]

    be = [
        (
            f"Entidades JPA, repositorios y adapters del alcance {hu_range} "
            f"(sobre el esquema definido y migrado por Fullstack)",
            18,
            "Backend",
        ),
        (f"Use cases + puertos del sprint ({hu_range})", 20, "Backend"),
        (f"Tests H2/tenant del alcance {hu_range}", 16, "Backend"),
        ("Colección API / fixtures de datos", 12, "Backend"),
        ("Pair hexagonal con Fullstack y corrección PR", 10, "Backend"),
        ("Ajustes por feedback QA", 4, "Backend"),
    ]

    fe = [
        (f"Pantallas/flujos Angular del alcance {hu_range}", 22, "Frontend"),
        ("Integración servicios domain/infra + estados UI", 16, "Frontend"),
        ("GrhUiTexts, vacíos, validaciones y permisos de sección", 14, "Frontend"),
        ("Pair UI con Fullstack", 12, "Frontend"),
        ("Ajustes UX por feedback QA", 10, "Frontend"),
        ("Regresión visual del sprint", 6, "Frontend"),
    ]

    qa = [
        (f"Diseño de casos y matriz de trazabilidad {hu_range}", 16, "QA"),
        ("Ejecución pruebas funcionales del sprint", 18, "QA"),
        ("Datos de prueba (empresa/frente/roles) y escenarios negativos", 12, "QA"),
        ("Retest defectos + regresión humo de sprints previos", 14, "QA"),
        ("Evidencia y checklist de cierre", 10, "QA"),
        ("Pair con Fullstack sobre criterios de aceptación", 10, "QA"),
    ]
    return fs + be + fe + qa


SPRINTS = [
    {
        "n": 1,
        "name": "Sprint 1",
        "title": "Integración GRH y catálogos núcleo (EP-00 / EP-01 inicio)",
        "objetivo": (
            "Habilitar el módulo en la plataforma (menú/submódulos Super Admin, permisos, alcance por frente, "
            "tenant, maestros GRH) y parametrizar catálogos núcleo: frentes, capacidades, turnos, horarios, "
            "pausas, estados, campañas, territorio, modalidades y sitios. "
            "La alta de módulo/submódulos en Super Admin es una actividad corta dentro de este sprint."
        ),
        "hu": "HU01–HU18",
        "hu_list": "HU01, HU02, HU03, HU04, HU05, HU06, HU07, HU08, HU09, HU10, HU11, HU12, HU13, HU14, HU15, HU16, HU17, HU18",
        "focus": "Integración GRH + Catálogos núcleo",
    },
    {
        "n": 2,
        "name": "Sprint 2",
        "title": "Catálogos avanzados y motor de reglas (EP-01 cierre)",
        "objetivo": (
            "Restricciones, festivos, cortes, tipos de hora, compensatorio, cobertura, "
            "motor de reglas, flags de atributos e importación asistida desde Excel."
        ),
        "hu": "HU19–HU27",
        "hu_list": "HU19, HU20, HU21, HU22, HU23, HU24, HU25, HU26, HU27",
        "focus": "Catálogos avanzados + motor",
    },
    {
        "n": 3,
        "name": "Sprint 3",
        "title": "Construcción de malla core (EP-02)",
        "objetivo": (
            "Crear malla, grupo, asignar turno/estado/extra/territorio/modalidad-sitio/nota "
            "y visualizar la grilla operativa."
        ),
        "hu": "HU28–HU35",
        "hu_list": "HU28, HU29, HU30, HU31, HU32, HU33, HU34, HU35",
        "focus": "Construcción de malla core",
    },
    {
        "n": 4,
        "name": "Sprint 4",
        "title": "Construcción avanzada (EP-02 cierre)",
        "objetivo": (
            "Campaña en celda, cobertura cruzada, filtros, concurrencia, carga parcial, "
            "copia masiva y fijar atributo por periodo."
        ),
        "hu": "HU36–HU42",
        "hu_list": "HU36, HU37, HU38, HU39, HU40, HU41, HU42",
        "focus": "Construcción avanzada",
    },
    {
        "n": 5,
        "name": "Sprint 5",
        "title": "Validación y rotación base (EP-03)",
        "objetivo": (
            "Contadores, equilibrio, conflictos, patrón, vínculo, simulación, "
            "exclusiones y aplicar rotación."
        ),
        "hu": "HU43–HU50",
        "hu_list": "HU43, HU44, HU45, HU46, HU47, HU48, HU49, HU50",
        "focus": "Validación y rotación base",
    },
    {
        "n": 6,
        "name": "Sprint 6",
        "title": "Rotación avanzada y publicación (EP-03 / EP-04)",
        "objetivo": (
            "Asistida, breaks, sitios, reequilibrio, restricciones al rotar, ajuste manual "
            "y ciclo de publicación."
        ),
        "hu": "HU51–HU59",
        "hu_list": "HU51, HU52, HU53, HU54, HU55, HU56, HU57, HU58, HU59",
        "focus": "Rotación avanzada + publicación",
    },
    {
        "n": 7,
        "name": "Sprint 7",
        "title": "Novedades, consulta y empleado (EP-05 / EP-06 / EP-07)",
        "objetivo": (
            "Novedades, ownership, historial, notificaciones, consulta operativa "
            "y vistas del empleado."
        ),
        "hu": "HU60–HU69",
        "hu_list": "HU60, HU61, HU62, HU63, HU64, HU65, HU66, HU67, HU68, HU69",
        "focus": "Novedades, consulta y empleado",
    },
    {
        "n": 8,
        "name": "Sprint 8",
        "title": "Reportes y horas (EP-08)",
        "objetivo": (
            "Cálculo de horas, export, novedad vs operativa, cobertura terceros, "
            "plantilla/import TH y compensatorio."
        ),
        "hu": "HU70–HU76",
        "hu_list": "HU70, HU71, HU72, HU73, HU74, HU75, HU76",
        "focus": "Reportes y horas",
    },
    {
        "n": 9,
        "name": "Sprint 9",
        "title": "Intercambio y estabilización (EP-09)",
        "objetivo": (
            "Configurar, solicitar, validar, aprobar y aplicar intercambio; "
            "hardening y salida a piloto."
        ),
        "hu": "HU77–HU81",
        "hu_list": "HU77, HU78, HU79, HU80, HU81",
        "focus": "Intercambio + estabilización",
    },
]


def _assert_hours(rows: list[tuple[str, int, str]]) -> None:
    totals: dict[str, int] = {}
    for _, h, role in rows:
        totals[role] = totals.get(role, 0) + h
    for role in ("Fullstack", "Backend", "Frontend", "QA"):
        if totals.get(role, 0) != CAP:
            raise SystemExit(f"{role} = {totals.get(role)} (esperado {CAP})")


def write_xlsx() -> None:
    wb = Workbook()
    # Resumen
    ws0 = wb.active
    ws0.title = "Resumen"
    header_fill = PatternFill("solid", fgColor="1B3A4B")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="D0D8DE"),
        right=Side(style="thin", color="D0D8DE"),
        top=Side(style="thin", color="D0D8DE"),
        bottom=Side(style="thin", color="D0D8DE"),
    )
    alt = PatternFill("solid", fgColor="F4F7F9")

    ws0["A1"] = "Backlog de sprints — Malla de Turnos"
    ws0["A1"].font = Font(bold=True, size=14, color="1B3A4B")
    ws0.merge_cells("A1:E1")
    meta = [
        ("Versión", "4.0 — Sprint 1…9 (sin Sprint 0); EP-00 en Sprint 1"),
        ("Integrantes", "1 Fullstack · 1 Backend · 1 Frontend · 1 QA"),
        ("Duración sprint", "2 semanas"),
        ("Capacidad / integrante", "80 h (equipo 320 h)"),
        ("Total sprints", "9 (Sprint 1 … Sprint 9)"),
        ("Duración orientativa", "9 × 2 semanas ≈ 18 semanas"),
        ("Nota", "Cada hoja = un sprint. Columnas listas para filtrar/ordenar en Excel."),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        ws0[f"A{i}"] = k
        ws0[f"A{i}"].font = Font(bold=True)
        ws0[f"B{i}"] = v
        ws0.merge_cells(f"B{i}:E{i}")

    ws0["A11"] = "SPRINT"
    ws0["B11"] = "ALCANCE HU"
    ws0["C11"] = "ENFOQUE"
    ws0["D11"] = "HORAS / INTEGRANTE"
    ws0["E11"] = "HORAS EQUIPO"
    for col in range(1, 6):
        cell = ws0.cell(11, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin

    for i, sp in enumerate(SPRINTS):
        r = 12 + i
        ws0.cell(r, 1, sp["name"]).border = thin
        ws0.cell(r, 2, sp["hu"]).border = thin
        ws0.cell(r, 3, sp["title"]).border = thin
        ws0.cell(r, 4, CAP).border = thin
        ws0.cell(r, 5, CAP * 4).border = thin
        if i % 2:
            for c in range(1, 6):
                ws0.cell(r, c).fill = alt

    ws0.column_dimensions["A"].width = 14
    ws0.column_dimensions["B"].width = 14
    ws0.column_dimensions["C"].width = 55
    ws0.column_dimensions["D"].width = 20
    ws0.column_dimensions["E"].width = 16

    # Hojas por sprint
    for sp in SPRINTS:
        rows = activities(sp["n"], sp["hu"], sp["focus"])
        _assert_hours(rows)
        ws = wb.create_sheet(sp["name"])

        ws["A1"] = sp["name"]
        ws["A1"].font = Font(bold=True, size=14, color="1B3A4B")
        ws["B1"] = sp["title"]
        ws["B1"].font = Font(bold=True, size=12, color="2E6B8A")
        ws.merge_cells("B1:C1")

        ws["A2"] = "Objetivo"
        ws["A2"].font = Font(bold=True)
        ws["B2"] = sp["objetivo"]
        ws.merge_cells("B2:C2")
        ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 48

        ws["A3"] = "Alcance HU"
        ws["A3"].font = Font(bold=True)
        ws["B3"] = sp["hu_list"]
        ws.merge_cells("B3:C3")
        ws["B3"].alignment = Alignment(wrap_text=True)

        # Encabezados de tabla (fila 5) — listos para copiar/filtrar
        headers = ["ACTIVIDADES", "HORAS PLANEADAS", "ROL"]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(5, c, h)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin
            cell.alignment = Alignment(horizontal="center")

        for i, (act, hours, role) in enumerate(rows):
            r = 6 + i
            ws.cell(r, 1, act).border = thin
            ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
            cell_h = ws.cell(r, 2, hours)
            cell_h.border = thin
            cell_h.alignment = Alignment(horizontal="center")
            cell_r = ws.cell(r, 3, role)
            cell_r.border = thin
            cell_r.alignment = Alignment(horizontal="center")
            if i % 2:
                for c in range(1, 4):
                    ws.cell(r, c).fill = alt

        # Totales por rol
        start = 6 + len(rows) + 1
        ws.cell(start, 1, "SUBTOTAL POR ROL").font = Font(bold=True)
        for j, role in enumerate(("Fullstack", "Backend", "Frontend", "QA")):
            ws.cell(start + 1 + j, 1, role)
            ws.cell(start + 1 + j, 2, CAP)
            ws.cell(start + 1 + j, 3, "OK")

        ws.column_dimensions["A"].width = 95
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 14
        ws.auto_filter.ref = f"A5:C{5 + len(rows)}"
        ws.freeze_panes = "A6"

    wb.save(XLSX)
    print(f"Wrote {XLSX}")


def write_md() -> None:
    lines = [
        "# Backlog de sprints — Malla de Turnos",
        "",
        "| Campo | Valor |",
        "| ----- | ----- |",
        "| Versión | 4.0 — Sprint 1…9 (EP-00 dentro de Sprint 1) |",
        "| Integrantes | 1 Fullstack · 1 Backend (aprendiz) · 1 Frontend (aprendiz) · 1 QA (aprendiz) |",
        "| Duración sprint | 2 semanas |",
        "| Capacidad por integrante | **80 h / sprint** (4 × 80 = **320 h** equipo) |",
        "| Total sprints | **9** (Sprint 1 … Sprint 9) |",
        "| Duración orientativa | 9 × 2 semanas ≈ **18 semanas** |",
        "| Orden | HU01→HU81 (EP-00→EP-09) |",
        "| Excel | `docs/BACKLOG-SPRINTS-MALLA-TURNOS.xlsx` (una hoja por sprint) |",
        "",
        "Verificación: en cada sprint Fullstack = 80, Backend = 80, Frontend = 80, QA = 80.",
        "",
        "Fullstack hace **diseño de BD solo en Sprint 1**; en sprints siguientes solo migraciones Flyway si el esquema crece. Backend **no** hace Flyway.",
        "",
        "---",
        "",
        "## Resumen",
        "",
        "| Sprint | Alcance HU | Enfoque | Horas/integrante |",
        "| ------ | ---------- | ------- | ---------------- |",
    ]
    for sp in SPRINTS:
        lines.append(f"| {sp['n']} | {sp['hu']} | {sp['title']} | 80 |")
    lines += ["", "---", ""]

    for sp in SPRINTS:
        rows = activities(sp["n"], sp["hu"], sp["focus"])
        _assert_hours(rows)
        lines += [
            f"## {sp['name']} — {sp['title']}",
            "",
            f"Objetivo: {sp['objetivo']}",
            "",
            f"Alcance HU: **{sp['hu_list']}**",
            "",
            "ACTIVIDADES\tHORAS PLANEADAS\tROL",
        ]
        for act, hours, role in rows:
            lines.append(f"{act}\t{hours}\t{role}")
        lines += ["", "Subtotal: Fullstack 80 · Backend 80 · Frontend 80 · QA 80", "", "---", ""]

    lines += [
        "## Criterios de priorización si hay que cortar",
        "",
        "1. Mantener: Sprint 1–4 y 6–8 (módulo usable: integra, parametriza, arma, publica, empleado ve, exporta horas).",
        "2. Diferir parcial Sprint 5–6 modo automático completo (dejar manual + conflictos/contadores).",
        "3. Diferir Sprint 9 (intercambio) si el frente piloto nace con la capacidad off y negocio lo acepta temporalmente.",
        "4. Fullstack: diseño BD en Sprint 1; Flyway siempre a cargo de Fullstack. Backend: use cases/JPA/adapters/tests, sin Flyway.",
        "",
    ]
    MD.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
    print(f"Wrote {MD}")


if __name__ == "__main__":
    write_xlsx()
    write_md()
