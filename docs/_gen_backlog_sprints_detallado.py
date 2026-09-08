# -*- coding: utf-8 -*-
"""
Propuesta DETALLADA de sprints (actividades específicas por rol).
NO modifica BACKLOG-SPRINTS-MALLA-TURNOS.xlsx / .md.

Reglas:
- Diseño de BD (modelo) solo en Sprint 1 → Fullstack.
- Flyway siempre Fullstack (nunca Backend).
- Backend: JPA/repos/adapters/use cases/tests sobre esquema ya migrado.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from _sprints_detallado_data import SPRINTS

DOCS = Path(__file__).resolve().parent
XLSX = DOCS / "BACKLOG-SPRINTS-MALLA-TURNOS-DETALLADO.xlsx"
MD = DOCS / "BACKLOG-SPRINTS-MALLA-TURNOS-DETALLADO.md"
CAP = 80
HU_REF = re.compile(r"HU\d+", re.IGNORECASE)

HEADER_FILL = PatternFill("solid", fgColor="1B3A4B")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT = PatternFill("solid", fgColor="F4F7F9")
THIN = Border(
    left=Side(style="thin", color="D0D8DE"),
    right=Side(style="thin", color="D0D8DE"),
    top=Side(style="thin", color="D0D8DE"),
    bottom=Side(style="thin", color="D0D8DE"),
)


def re_flyway(act: str) -> bool:
    return "flyway" in act.lower()


def check(rows: list[tuple[str, int, str]], sprint: str) -> None:
    totals: dict[str, int] = {}
    for act, h, r in rows:
        totals[r] = totals.get(r, 0) + h
        if r == "Backend" and re_flyway(act):
            raise SystemExit(f"{sprint}: Backend no debe tener Flyway → {act[:60]}")
        if r == "Fullstack" and "Diseño BD" in act and sprint != "Sprint 1":
            raise SystemExit(f"{sprint}: Diseño BD solo en Sprint 1 → {act[:60]}")
        if not HU_REF.search(act):
            raise SystemExit(f"{sprint}: actividad sin referencia HU → {act[:80]}")
    for role in ("Fullstack", "Backend", "Frontend", "QA"):
        if totals.get(role, 0) != CAP:
            raise SystemExit(f"{sprint} {role}={totals.get(role)} (esperado {CAP})")


def write_xlsx() -> None:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Resumen"
    ws0["A1"] = "Propuesta DETALLADA de sprints — Malla de Turnos"
    ws0["A1"].font = Font(bold=True, size=14, color="1B3A4B")
    ws0.merge_cells("A1:E1")
    meta = [
        ("Archivo", "NUEVO — no reemplaza BACKLOG-SPRINTS-MALLA-TURNOS.xlsx"),
        ("Versión", "1.2 detallada — cada actividad con HU; BD solo Sprint 1; Backend sin Flyway"),
        ("Sprints", "9 (Sprint 1 … Sprint 9); EP-00 dentro de Sprint 1"),
        ("Capacidad", "80 h / rol / sprint (equipo 320 h)"),
        ("Fullstack", "Diseño BD solo Sprint 1; Flyway siempre Fullstack"),
        ("Backend", "JPA/repos/adapters/use cases/tests — sin Flyway"),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        ws0[f"A{i}"] = k
        ws0[f"A{i}"].font = Font(bold=True)
        ws0[f"B{i}"] = v
        ws0.merge_cells(f"B{i}:E{i}")

    headers = ["SPRINT", "ALCANCE HU", "ENFOQUE", "HORAS/INTEGRANTE", "HORAS EQUIPO"]
    for c, h in enumerate(headers, 1):
        cell = ws0.cell(11, c, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN

    for i, sp in enumerate(SPRINTS):
        check(sp["rows"], sp["name"])
        r = 12 + i
        vals = [sp["name"], sp["hu"], sp["title"], CAP, CAP * 4]
        for c, v in enumerate(vals, 1):
            cell = ws0.cell(r, c, v)
            cell.border = THIN
            if i % 2:
                cell.fill = ALT

    ws0.column_dimensions["A"].width = 12
    ws0.column_dimensions["B"].width = 14
    ws0.column_dimensions["C"].width = 58
    ws0.column_dimensions["D"].width = 18
    ws0.column_dimensions["E"].width = 14

    for sp in SPRINTS:
        rows = sp["rows"]
        check(rows, sp["name"])
        ws = wb.create_sheet(sp["name"])
        ws["A1"] = sp["name"]
        ws["A1"].font = Font(bold=True, size=14, color="1B3A4B")
        ws["B1"] = sp["title"]
        ws["B1"].font = Font(bold=True, size=12, color="2E6B8A")
        ws.merge_cells("B1:C1")
        ws["A2"] = "Objetivo"
        ws["A2"].font = Font(bold=True)
        ws["B2"] = sp["objetivo"]
        ws.merge_cells("B2:C2")
        ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 45
        ws["A3"] = "Alcance HU"
        ws["A3"].font = Font(bold=True)
        ws["B3"] = sp["hu_list"]
        ws.merge_cells("B3:C3")

        for c, h in enumerate(["ACTIVIDADES", "HORAS PLANEADAS", "ROL"], 1):
            cell = ws.cell(5, c, h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = THIN

        for i, (act, hours, role) in enumerate(rows):
            r = 6 + i
            ws.cell(r, 1, act).border = THIN
            ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, 2, hours).border = THIN
            ws.cell(r, 2).alignment = Alignment(horizontal="center")
            ws.cell(r, 3, role).border = THIN
            ws.cell(r, 3).alignment = Alignment(horizontal="center")
            if i % 2:
                for c in range(1, 4):
                    ws.cell(r, c).fill = ALT

        start = 6 + len(rows) + 1
        ws.cell(start, 1, "SUBTOTAL POR ROL").font = Font(bold=True)
        for j, role in enumerate(("Fullstack", "Backend", "Frontend", "QA")):
            ws.cell(start + 1 + j, 1, role)
            ws.cell(start + 1 + j, 2, CAP)
            ws.cell(start + 1 + j, 3, "OK")

        ws.column_dimensions["A"].width = 100
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 12
        ws.auto_filter.ref = f"A5:C{5 + len(rows)}"
        ws.freeze_panes = "A6"

    wb.save(XLSX)
    print(f"Wrote {XLSX}")


def write_md() -> None:
    lines = [
        "# Backlog de sprints DETALLADO — Malla de Turnos",
        "",
        "> Propuesta nueva. **No reemplaza** `BACKLOG-SPRINTS-MALLA-TURNOS.xlsx`.",
        "",
        "| Campo | Valor |",
        "| ----- | ----- |",
        "| Archivo Excel | `docs/BACKLOG-SPRINTS-MALLA-TURNOS-DETALLADO.xlsx` |",
        "| Versión | 1.2 — cada actividad con HU; BD solo Sprint 1; Backend sin Flyway |",
        "| Sprints | 9 (Sprint 1 … Sprint 9) |",
        "| Capacidad | 80 h / rol / sprint |",
        "",
        "- **Diseño de BD (modelo):** solo Sprint 1, rol Fullstack.",
        "- **Flyway:** siempre Fullstack.",
        "- **Backend:** JPA/repos/adapters/use cases/tests; no Flyway.",
        "- **HU en actividades:** cada fila indica entre paréntesis la(s) HU que cubre.",
        "",
        "---",
        "",
    ]
    for sp in SPRINTS:
        check(sp["rows"], sp["name"])
        lines += [
            f"## {sp['name']} — {sp['title']}",
            "",
            f"Objetivo: {sp['objetivo']}",
            "",
            f"Alcance HU: **{sp['hu_list']}**",
            "",
            "ACTIVIDADES\tHORAS PLANEADAS\tROL",
        ]
        for act, hours, role in sp["rows"]:
            lines.append(f"{act}\t{hours}\t{role}")
        lines += ["", "Subtotal: Fullstack 80 · Backend 80 · Frontend 80 · QA 80", "", "---", ""]
    MD.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
    print(f"Wrote {MD}")


if __name__ == "__main__":
    write_xlsx()
    write_md()
