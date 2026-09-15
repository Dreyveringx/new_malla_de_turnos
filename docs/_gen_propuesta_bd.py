# -*- coding: utf-8 -*-
"""Genera PROPUESTA-BD-MALLA-TURNOS.xlsx / .csv / .dbml (post sesion 11-sep-2026)."""
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    import subprocess
    import sys

    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).resolve().parent
rows = []


def add(grupo, table, col, typ, desc):
    rows.append((grupo, table, col, typ, desc))


G1 = "1. Catalogos / Configuracion"
G2 = "2. Reglas internas y horas"
G3 = "3. Operacion / Planificacion"
G4 = "4. Intercambio, novedades e importacion"
G5 = "5. NO crear en ms_malla_turnos (externo GRH)"

# --- operational_front ---
T = "operational_front"
add(G1, T, "id", "UUID PK", "Identificador unico del frente operativo.")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant GRH (ms_company_admin). Sin FK cross-DB.")
add(G1, T, "code", "VARCHAR(20) NOT NULL", "Codigo del frente. UNIQUE(company_id, code). Preferible auto-consecutivo (#3).")
add(G1, T, "name", "VARCHAR(120) NOT NULL", "Nombre visible (ej. Contact Center).")
add(G1, T, "description", "TEXT NULL", "Descripcion de ayuda.")
add(G1, T, "builder_profile_label", "VARCHAR(120) NULL", "Perfil/rol de negocio que arma la malla (texto guia).")
add(G1, T, "publisher_profile_label", "VARCHAR(120) NULL", "Perfil/rol de negocio que publica/aprueba.")
add(G1, T, "separate_builder_publisher", "BOOLEAN NOT NULL DEFAULT FALSE", "Si true, quien arma no publica (ciclo revision).")
add(G1, T, "build_mode", "VARCHAR(20) NOT NULL", "manual | asistido | automatico.")
add(G1, T, "period_type", "VARCHAR(10) NOT NULL", "semana | quincena | mes (Mesa semanal; CC quincenal).")
add(G1, T, "published_editable", "BOOLEAN NOT NULL DEFAULT FALSE", "Permite reasignar celdas tras publicar (con historial).")
add(G1, T, "swap_enabled", "BOOLEAN NOT NULL DEFAULT FALSE", "Intercambio ON/OFF. Sin tope mensual de solicitudes (#16).")
add(G1, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "Inactivo no aparece en nuevas mallas.")
add(G1, T, "created_at", "TIMESTAMPTZ NOT NULL", "Alta.")
add(G1, T, "updated_at", "TIMESTAMPTZ NOT NULL", "Ultima modificacion.")
add(G1, T, "created_by", "UUID NULL", "user_id logico ms_auth.")
add(G1, T, "updated_by", "UUID NULL", "user_id logico ms_auth.")

T = "front_capability"
add(G1, T, "id", "UUID PK", "PK.")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G1, T, "capability_code", "VARCHAR(40) NOT NULL", "usa_territorio | usa_campanas | usa_modalidad | usa_sitio | usa_multi_break | usa_multi_zona")
add(G1, T, "enabled", "BOOLEAN NOT NULL DEFAULT FALSE", "Si la grilla/asignacion muestra el atributo.")
add(G1, T, "config_json", "JSONB NULL", "Parametros opcionales del flag.")
add(G1, T, "UNIQUE", "(company_id, front_id, capability_code)", "Una capacidad por frente.")

T = "front_area_link"
add(G1, T, "id", "UUID PK", "PK.")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G1, T, "area_id", "UUID NOT NULL", "UUID logico de area en ms_parametrization (sin FK).")
add(G1, T, "UNIQUE", "(company_id, front_id, area_id)", "Evita duplicar vinculo.")

T = "front_user_scope"
add(G1, T, "id", "UUID PK", "Alcance de frentes por usuario (HU03).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G1, T, "user_id", "UUID NOT NULL", "UUID logico usuario ms_auth.")
add(G1, T, "UNIQUE", "(company_id, front_id, user_id)", "")

T = "front_role_scope"
add(G1, T, "id", "UUID PK", "Alcance de frentes por rol (HU03).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G1, T, "role_id", "UUID NOT NULL", "UUID logico rol ms_auth.")
add(G1, T, "UNIQUE", "(company_id, front_id, role_id)", "")

T = "catalog_code_sequence"
add(G1, T, "id", "UUID PK", "Secuencia de codigos automaticos por empresa y tipo (#3).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "catalog_type", "VARCHAR(40) NOT NULL", "operational_front | shift_template | cell_state | work_modality | campaign | territory_node")
add(G1, T, "prefix", "VARCHAR(10) NULL", "Prefijo opcional (ej. T, EST).")
add(G1, T, "next_value", "INT NOT NULL DEFAULT 1", "Siguiente consecutivo.")
add(G1, T, "UNIQUE", "(company_id, catalog_type)", "Una secuencia por tipo.")

T = "shift_template"
add(G1, T, "id", "UUID PK", "Plantilla de turno (catalogo unico empresa; se habilita a frentes).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "code", "VARCHAR(20) NOT NULL", "UNIQUE(company_id, code). Auto-consecutivo recomendado.")
add(G1, T, "name", "VARCHAR(120) NOT NULL", "Nombre visible (ej. Manana A).")
add(G1, T, "start_time", "TIME NOT NULL", "Hora inicio base. Un turno conserva su horario sin importar la zona.")
add(G1, T, "end_time", "TIME NOT NULL", "Hora fin base.")
add(G1, T, "crosses_midnight", "BOOLEAN NOT NULL DEFAULT FALSE", "Fin al dia siguiente.")
add(G1, T, "color_hex", "VARCHAR(7) NULL", "Color en grilla (#RRGGBB).")
add(G1, T, "net_hours", "NUMERIC(5,2) NULL", "Horas netas de referencia (opcional).")
add(G1, T, "multi_break_enabled", "BOOLEAN NOT NULL DEFAULT FALSE", "Permite varios breaks en el turno (#11).")
add(G1, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "Disponible para asignar.")
add(G1, T, "created_at", "TIMESTAMPTZ NOT NULL", "")
add(G1, T, "updated_at", "TIMESTAMPTZ NOT NULL", "")

T = "shift_template_front"
add(G1, T, "id", "UUID PK", "N:M turno <-> frentes (habilitacion). UI: dropdown multi (#1).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "shift_template_id", "UUID NOT NULL", "FK -> shift_template.id")
add(G1, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G1, T, "UNIQUE", "(company_id, shift_template_id, front_id)", "")

T = "shift_template_day_schedule"
add(G1, T, "id", "UUID PK", "Horario distinto por dia de semana (HU12).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "shift_template_id", "UUID NOT NULL", "FK -> shift_template.id")
add(G1, T, "dow", "SMALLINT NOT NULL", "Dia de semana (0=domingo .. 6=sabado o 1-7 segun convencion).")
add(G1, T, "start_time", "TIME NOT NULL", "Inicio ese dia.")
add(G1, T, "end_time", "TIME NOT NULL", "Fin ese dia.")
add(G1, T, "crosses_midnight", "BOOLEAN NOT NULL DEFAULT FALSE", "")
add(G1, T, "UNIQUE", "(shift_template_id, dow)", "Un horario por dia.")

T = "shift_template_break"
add(G1, T, "id", "UUID PK", "Break/almuerzo de la plantilla. Varios si multi_break_enabled (#11).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "shift_template_id", "UUID NOT NULL", "FK -> shift_template.id")
add(G1, T, "code", "VARCHAR(20) NULL", "Codigo interno del break (opcional).")
add(G1, T, "label", "VARCHAR(80) NOT NULL", "Descanso | Refrigerio | Almuerzo | especial")
add(G1, T, "duration_minutes", "INT NOT NULL", "Duracion en minutos.")
add(G1, T, "paid", "BOOLEAN NOT NULL DEFAULT FALSE", "Si suma como tiempo pagado.")
add(G1, T, "sort_order", "SMALLINT NOT NULL DEFAULT 0", "Orden de presentacion.")
add(G1, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "")

T = "cell_state"
add(G1, T, "id", "UUID PK", "Estado de celda (DES, VAC, INC, PER, CAP...).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant. Catalogo de empresa.")
add(G1, T, "code", "VARCHAR(20) NOT NULL", "UNIQUE(company_id, code).")
add(G1, T, "name", "VARCHAR(120) NOT NULL", "Nombre visible.")
add(G1, T, "color_hex", "VARCHAR(7) NULL", "Color en grilla.")
add(G1, T, "assignable", "BOOLEAN NOT NULL DEFAULT TRUE", "Se puede asignar en armado.")
add(G1, T, "counts_hours", "BOOLEAN NOT NULL DEFAULT FALSE", "Suma horas productivas/reporte.")
add(G1, T, "deducts_productive", "BOOLEAN NOT NULL DEFAULT FALSE", "Descuenta de productividad.")
add(G1, T, "requires_support", "BOOLEAN NOT NULL DEFAULT FALSE", "Requiere soporte/evidencia.")
add(G1, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "")

T = "work_modality"
add(G1, T, "id", "UUID PK", "Modalidad GLOBAL de empresa (#2: SIN vinculo a frente).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "code", "VARCHAR(20) NOT NULL", "UNIQUE(company_id, code). Auto.")
add(G1, T, "name", "VARCHAR(120) NOT NULL", "Presencial | Virtual | Hibrido")
add(G1, T, "icon", "VARCHAR(40) NULL", "Icono UI opcional.")
add(G1, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "Se elige al asignar persona/celda.")

T = "attendance_site"
add(G1, T, "id", "UUID PK", "Sitio de asistencia/sede.")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "code", "VARCHAR(20) NOT NULL", "UNIQUE(company_id, code). Puede ser MANUAL codigo de negocio (#3).")
add(G1, T, "name", "VARCHAR(120) NOT NULL", "Nombre del sitio.")
add(G1, T, "site_type", "VARCHAR(40) NULL", "Sede | SPT | Laboratorio | Remoto")
add(G1, T, "address", "VARCHAR(255) NULL", "Direccion o referencia.")
add(G1, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "")

T = "attendance_site_front"
add(G1, T, "id", "UUID PK", "Sitios aplicables por frente (opcional).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "site_id", "UUID NOT NULL", "FK -> attendance_site.id")
add(G1, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G1, T, "UNIQUE", "(company_id, site_id, front_id)", "")

T = "campaign"
add(G1, T, "id", "UUID PK", "Campana/tarea informativa. NO va en paso 1 de crear malla (#17).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "code", "VARCHAR(20) NOT NULL", "UNIQUE(company_id, code). Auto.")
add(G1, T, "name", "VARCHAR(120) NOT NULL", "Nombre.")
add(G1, T, "front_id", "UUID NULL", "Si NULL aplica amplio; si se setea, acota al frente.")
add(G1, T, "valid_from", "DATE NULL", "Vigencia inicio.")
add(G1, T, "valid_to", "DATE NULL", "Vigencia fin.")
add(G1, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "")

T = "territory_level"
add(G1, T, "id", "UUID PK", "Nivel jerarquico configurable (Regional / Zona / SPT).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "level_order", "SMALLINT NOT NULL", "1, 2, 3...")
add(G1, T, "name", "VARCHAR(80) NOT NULL", "Nombre del nivel.")
add(G1, T, "UNIQUE", "(company_id, level_order)", "")

T = "territory_node"
add(G1, T, "id", "UUID PK", "Nodo de territorio (ej. SPT Soacha).")
add(G1, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G1, T, "level_id", "UUID NOT NULL", "FK -> territory_level.id")
add(G1, T, "parent_id", "UUID NULL", "FK -> territory_node.id (NULL = raiz).")
add(G1, T, "code", "VARCHAR(20) NOT NULL", "UNIQUE(company_id, code). Auto recomendado.")
add(G1, T, "name", "VARCHAR(120) NOT NULL", "Nombre del nodo.")
add(G1, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "")

# --- rules ---
T = "validation_rule"
add(G2, T, "id", "UUID PK", "Reglas internas del motor. Sin pantalla de cobertura/compensatorio (#4 #5).")
add(G2, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G2, T, "code", "VARCHAR(40) NOT NULL", "NO_SOLAPE | MAX_HORAS_PERIODO | MAX_DIAS_CONSECUTIVOS")
add(G2, T, "scope", "VARCHAR(20) NOT NULL", "empresa | frente")
add(G2, T, "front_id", "UUID NULL", "FK -> operational_front si scope=frente.")
add(G2, T, "param_value", "VARCHAR(40) NULL", "Valor (ej. 46). Seed/admin tecnico, no UI de negocio.")
add(G2, T, "param_unit", "VARCHAR(40) NULL", "horas | dias")
add(G2, T, "severity", "VARCHAR(20) NOT NULL", "info | advertencia | bloqueo")
add(G2, T, "priority", "INT NOT NULL DEFAULT 100", "Orden de evaluacion.")
add(G2, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "")
add(G2, T, "notes", "TEXT NULL", "Documentacion interna.")

T = "hour_classification_band"
add(G2, T, "id", "UUID PK", "Franjas de recargo/clasificacion de horas (#8). Definir con Administracion.")
add(G2, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G2, T, "code", "VARCHAR(20) NOT NULL", "ORDI | NOCT | DOM | FEST | HEXA | HEXN. UNIQUE(company_id, code)")
add(G2, T, "name", "VARCHAR(120) NOT NULL", "Nombre para reporteria.")
add(G2, T, "band_type", "VARCHAR(40) NOT NULL", "ordinaria | nocturna | dominical | festiva | extra_diurna | extra_nocturna")
add(G2, T, "start_time", "TIME NULL", "Inicio franja (si aplica por hora).")
add(G2, T, "end_time", "TIME NULL", "Fin franja.")
add(G2, T, "crosses_midnight", "BOOLEAN NOT NULL DEFAULT FALSE", "")
add(G2, T, "applies_sunday", "BOOLEAN NOT NULL DEFAULT FALSE", "")
add(G2, T, "applies_holiday", "BOOLEAN NOT NULL DEFAULT FALSE", "Festivos desde ms_parametrization (no tabla local).")
add(G2, T, "sort_order", "SMALLINT NOT NULL DEFAULT 0", "")
add(G2, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "")

T = "rotation_pattern"
add(G2, T, "id", "UUID PK", "Patron de rotacion. Al cerrar periodo se aplica; excepciones manuales (#12).")
add(G2, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G2, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G2, T, "code", "VARCHAR(20) NOT NULL", "")
add(G2, T, "name", "VARCHAR(120) NOT NULL", "")
add(G2, T, "cycle_length", "SMALLINT NOT NULL", "Largo del ciclo (pasos).")
add(G2, T, "auto_apply_on_period_close", "BOOLEAN NOT NULL DEFAULT TRUE", "Automatizar al cerrar periodo (#12).")
add(G2, T, "active", "BOOLEAN NOT NULL DEFAULT TRUE", "")

T = "rotation_pattern_step"
add(G2, T, "id", "UUID PK", "Paso del patron (secuencia).")
add(G2, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G2, T, "pattern_id", "UUID NOT NULL", "FK -> rotation_pattern.id")
add(G2, T, "step_order", "SMALLINT NOT NULL", "Orden 1..N")
add(G2, T, "shift_template_id", "UUID NULL", "FK -> shift_template.id")
add(G2, T, "cell_state_id", "UUID NULL", "FK -> cell_state.id si el paso es estado.")

# --- operation ---
T = "schedule_grid"
add(G3, T, "id", "UUID PK", "Cabecera de malla (periodo + frente).")
add(G3, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G3, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G3, T, "name", "VARCHAR(160) NULL", "Nombre visible (ej. CC Sep Q1 2026).")
add(G3, T, "period_start", "DATE NOT NULL", "Inicio del periodo.")
add(G3, T, "period_end", "DATE NOT NULL", "Fin del periodo.")
add(G3, T, "period_type", "VARCHAR(10) NOT NULL", "semana | quincena | mes (copia del frente al crear).")
add(G3, T, "status", "VARCHAR(20) NOT NULL", "borrador | revision | publicada | cerrada | rechazada")
add(G3, T, "build_mode", "VARCHAR(20) NOT NULL", "manual | asistido | automatico")
add(G3, T, "owner_user_id", "UUID NULL", "Quien construye (ms_auth).")
add(G3, T, "publisher_user_id", "UUID NULL", "Quien publica (ms_auth).")
add(G3, T, "rejection_reason", "TEXT NULL", "Motivo de rechazo.")
add(G3, T, "published_at", "TIMESTAMPTZ NULL", "Momento de publicacion.")
add(G3, T, "closed_at", "TIMESTAMPTZ NULL", "Cierre de periodo (puede disparar rotacion).")
add(G3, T, "version", "INT NOT NULL DEFAULT 1", "Optimistic lock / concurrencia.")
add(G3, T, "created_at", "TIMESTAMPTZ NOT NULL", "")
add(G3, T, "updated_at", "TIMESTAMPTZ NOT NULL", "")
add(G3, T, "NOTE", "sin campaign_id", "NO incluye campana en cabecera (#17). Campana va en celda.")

T = "schedule_grid_person"
add(G3, T, "id", "UUID PK", "Persona incluida en la malla (paso asignacion #9).")
add(G3, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G3, T, "grid_id", "UUID NOT NULL", "FK -> schedule_grid.id")
add(G3, T, "employee_id", "UUID NOT NULL", "UUID logico ms_employee (sin FK).")
add(G3, T, "sort_order", "INT NOT NULL DEFAULT 0", "Orden en grilla.")
add(G3, T, "modality_id", "UUID NULL", "FK -> work_modality.id (fijada por periodo si aplica).")
add(G3, T, "site_id", "UUID NULL", "FK -> attendance_site.id opcional por periodo.")
add(G3, T, "UNIQUE", "(grid_id, employee_id)", "Una fila por persona en la malla.")

T = "schedule_grid_person_territory"
add(G3, T, "id", "UUID PK", "Una o varias zonas/SPT por persona en el periodo (#10 #14).")
add(G3, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G3, T, "grid_person_id", "UUID NOT NULL", "FK -> schedule_grid_person.id")
add(G3, T, "territory_node_id", "UUID NOT NULL", "FK -> territory_node.id")
add(G3, T, "is_primary", "BOOLEAN NOT NULL DEFAULT FALSE", "Zona principal opcional.")
add(G3, T, "UNIQUE", "(grid_person_id, territory_node_id)", "")

T = "schedule_cell"
add(G3, T, "id", "UUID PK", "Celda persona x dia. Reasignar deja historial.")
add(G3, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G3, T, "grid_id", "UUID NOT NULL", "FK -> schedule_grid.id")
add(G3, T, "employee_id", "UUID NOT NULL", "UUID logico ms_employee.")
add(G3, T, "work_date", "DATE NOT NULL", "Dia calendario.")
add(G3, T, "slot", "SMALLINT NOT NULL DEFAULT 0", "0=principal; 1+=turno extra mismo dia.")
add(G3, T, "shift_template_id", "UUID NULL", "FK -> shift_template.id")
add(G3, T, "cell_state_id", "UUID NULL", "FK -> cell_state.id (XOR logico con turno).")
add(G3, T, "modality_id", "UUID NULL", "FK -> work_modality.id (override del dia).")
add(G3, T, "site_id", "UUID NULL", "FK -> attendance_site.id")
add(G3, T, "campaign_id", "UUID NULL", "FK -> campaign.id (atributo de celda).")
add(G3, T, "note", "TEXT NULL", "Observacion.")
add(G3, T, "locked_attrs", "JSONB NULL", "Atributos fijados por periodo (HU42).")
add(G3, T, "origin", "VARCHAR(40) NULL", "manual | rotacion | intercambio | import | novedad")
add(G3, T, "updated_at", "TIMESTAMPTZ NOT NULL", "")
add(G3, T, "updated_by", "UUID NULL", "user_id ms_auth.")
add(G3, T, "UNIQUE", "(grid_id, employee_id, work_date, slot)", "Celda logica unica.")

T = "schedule_cell_territory"
add(G3, T, "id", "UUID PK", "Zonas/SPT del dia (override multi-zona).")
add(G3, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G3, T, "cell_id", "UUID NOT NULL", "FK -> schedule_cell.id")
add(G3, T, "territory_node_id", "UUID NOT NULL", "FK -> territory_node.id")
add(G3, T, "UNIQUE", "(cell_id, territory_node_id)", "")

T = "schedule_cell_break"
add(G3, T, "id", "UUID PK", "Breaks de la celda. Permite repetir / especiales (#11).")
add(G3, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G3, T, "cell_id", "UUID NOT NULL", "FK -> schedule_cell.id")
add(G3, T, "template_break_id", "UUID NULL", "FK -> shift_template_break.id si viene de plantilla.")
add(G3, T, "label", "VARCHAR(80) NOT NULL", "Etiqueta del break.")
add(G3, T, "duration_minutes", "INT NOT NULL", "Minutos.")
add(G3, T, "paid", "BOOLEAN NOT NULL DEFAULT FALSE", "")
add(G3, T, "sort_order", "SMALLINT NOT NULL DEFAULT 0", "")

T = "schedule_cell_history"
add(G3, T, "id", "UUID PK", "Historial append-only (anterior/nuevo, motivo, usuario, fecha).")
add(G3, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G3, T, "cell_id", "UUID NOT NULL", "FK -> schedule_cell.id (sin ON DELETE CASCADE destructivo).")
add(G3, T, "grid_id", "UUID NOT NULL", "Desnormalizado para consultas.")
add(G3, T, "employee_id", "UUID NOT NULL", "Desnormalizado.")
add(G3, T, "work_date", "DATE NOT NULL", "Desnormalizado.")
add(G3, T, "action", "VARCHAR(40) NOT NULL", "reasignar | publicar | intercambio | novedad | import")
add(G3, T, "before_json", "JSONB NULL", "Snapshot anterior.")
add(G3, T, "after_json", "JSONB NULL", "Snapshot nuevo.")
add(G3, T, "reason", "TEXT NULL", "Justificacion.")
add(G3, T, "origin", "VARCHAR(40) NULL", "Origen del cambio.")
add(G3, T, "changed_by", "UUID NOT NULL", "user_id ms_auth.")
add(G3, T, "changed_at", "TIMESTAMPTZ NOT NULL", "")

# --- swap / novelty / import ---
T = "shift_swap_request"
add(G4, T, "id", "UUID PK", "Solicitud de cambio de turno. Sin limite mensual (#16).")
add(G4, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G4, T, "front_id", "UUID NOT NULL", "FK -> operational_front.id")
add(G4, T, "grid_id", "UUID NOT NULL", "FK -> schedule_grid.id")
add(G4, T, "requester_employee_id", "UUID NOT NULL", "ms_employee.")
add(G4, T, "counterpart_employee_id", "UUID NOT NULL", "ms_employee.")
add(G4, T, "requester_cell_id", "UUID NOT NULL", "FK -> schedule_cell.id")
add(G4, T, "counterpart_cell_id", "UUID NOT NULL", "FK -> schedule_cell.id")
add(G4, T, "status", "VARCHAR(30) NOT NULL", "pendiente_validacion | pendiente_aprobacion | aprobada | rechazada | aplicada | cancelada")
add(G4, T, "rule_warnings", "JSONB NULL", "Resultado motor de reglas.")
add(G4, T, "decision_reason", "TEXT NULL", "Motivo aprobacion/rechazo.")
add(G4, T, "decided_by", "UUID NULL", "user_id ms_auth.")
add(G4, T, "decided_at", "TIMESTAMPTZ NULL", "")
add(G4, T, "created_at", "TIMESTAMPTZ NOT NULL", "")
add(G4, T, "updated_at", "TIMESTAMPTZ NOT NULL", "")

T = "schedule_novelty"
add(G4, T, "id", "UUID PK", "Novedad operativa vinculada a la malla (#18).")
add(G4, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G4, T, "grid_id", "UUID NOT NULL", "FK -> schedule_grid.id")
add(G4, T, "cell_id", "UUID NULL", "FK -> schedule_cell.id si aplica a un dia.")
add(G4, T, "employee_id", "UUID NOT NULL", "ms_employee.")
add(G4, T, "work_date", "DATE NOT NULL", "Dia de la novedad.")
add(G4, T, "cell_state_id", "UUID NULL", "Estado aplicado (INC, PER...).")
add(G4, T, "description", "TEXT NULL", "Detalle.")
add(G4, T, "source", "VARCHAR(40) NOT NULL", "manual | th_import | sistema")
add(G4, T, "created_by", "UUID NULL", "ms_auth.")
add(G4, T, "created_at", "TIMESTAMPTZ NOT NULL", "")

T = "excel_import_batch"
add(G4, T, "id", "UUID PK", "Importacion asistida / plantilla Excel (#20).")
add(G4, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G4, T, "front_id", "UUID NULL", "Contexto opcional.")
add(G4, T, "import_type", "VARCHAR(40) NOT NULL", "catalogos | malla_borrador | novedades_th")
add(G4, T, "file_name", "VARCHAR(255) NULL", "Nombre archivo.")
add(G4, T, "status", "VARCHAR(20) NOT NULL", "cargado | descubierto | confirmado | aplicado | error")
add(G4, T, "summary_json", "JSONB NULL", "Conteo de filas / hallazgos.")
add(G4, T, "uploaded_by", "UUID NULL", "ms_auth.")
add(G4, T, "created_at", "TIMESTAMPTZ NOT NULL", "")

T = "excel_import_row"
add(G4, T, "id", "UUID PK", "Fila staging del Excel.")
add(G4, T, "company_id", "UUID NOT NULL", "Tenant.")
add(G4, T, "batch_id", "UUID NOT NULL", "FK -> excel_import_batch.id")
add(G4, T, "row_number", "INT NOT NULL", "Numero de fila.")
add(G4, T, "payload_json", "JSONB NOT NULL", "Datos crudos normalizados.")
add(G4, T, "action", "VARCHAR(40) NULL", "crear | mapear | ignorar")
add(G4, T, "status", "VARCHAR(20) NOT NULL", "pendiente | aceptada | rechazada | aplicada")
add(G4, T, "message", "TEXT NULL", "Error/aviso.")

# --- external ---
add(G5, "EXTERNO", "company_id", "UUID logico", "Empresa: ms_company_admin / JWT. Tenant en todas las tablas.")
add(G5, "EXTERNO", "employee_id", "UUID logico", "Empleado: ms_employee via API.")
add(G5, "EXTERNO", "area_id / cargo_id", "UUID logico", "ms_parametrization.")
add(G5, "EXTERNO", "user_id / role_id", "UUID logico", "ms_auth (permisos, alcance, auditoria).")
add(G5, "EXTERNO", "festivos", "API lectura", "Calendario empresa en ms_parametrization. NO tabla local de cortes (#6).")
add(G5, "EXTERNO", "notificaciones", "Event types", "ms_notification — publicacion y cambios (#13).")
add(G5, "ELIMINADO", "coverage_rule UI", "N/A", "Cobertura por turno/franjas; sin pantalla (#4).")
add(G5, "ELIMINADO", "compensatory_rule UI", "N/A", "Reglas internas (#5).")
add(G5, "ELIMINADO", "payroll_cutoff", "N/A", "Pantalla/calendario de cortes eliminados (#6).")
add(G5, "ELIMINADO", "work_modality_front", "N/A", "Modalidad global (#2).")
add(G5, "ELIMINADO", "swap_monthly_limit", "N/A", "Solicitudes ilimitadas (#16).")

# CSV
csv_path = OUT / "PROPUESTA-BD-MALLA-TURNOS-COLUMNAS.csv"
with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
    f.write("Grupo;Tabla;Columna;Tipo_dato;Descripcion\n")
    for g, t, c, ty, d in rows:
        def esc(x):
            return '"' + str(x).replace('"', '""') + '"'

        f.write(";".join(esc(x) for x in (g, t, c, ty, d)) + "\n")

# XLSX
wb = Workbook()
ws = wb.active
ws.title = "Columnas"
ws.append(["Grupo", "Tabla", "Columna", "Tipo_dato", "Descripcion"])
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E79")
for r in rows:
    ws.append(list(r))
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 28
ws.column_dimensions["D"].width = 42
ws.column_dimensions["E"].width = 95
ws.auto_filter.ref = ws.dimensions
ws.freeze_panes = "A2"

ws2 = wb.create_sheet("Tablas_resumen")
ws2.append(["Grupo", "Tabla", "Proposito", "Notas_sesion_11_09"])
for cell in ws2[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E79")
for r in [
    ("1. Catalogos", "operational_front", "Frente por empresa", "period_type con quincena; perfiles arma/publica"),
    ("1. Catalogos", "front_capability", "Flags por frente", "zona/breaks opcionales"),
    ("1. Catalogos", "shift_template (+front/day/break)", "Catalogo unico turnos", "horario fijo; zona no cambia turno"),
    ("1. Catalogos", "work_modality", "Modalidades globales", "#2 sin frentes"),
    ("1. Catalogos", "territory_*", "Regional/zona/SPT", "multi-zona persona"),
    ("2. Reglas", "validation_rule", "Motor interno", "#4 #5 sin pantallas cobertura"),
    ("2. Reglas", "hour_classification_band", "Franjas recargo", "#8 con Admin"),
    ("2. Reglas", "rotation_pattern", "Rotacion", "#12 auto al cerrar periodo"),
    ("3. Operacion", "schedule_grid", "Malla", "sin campana en cabecera #17"),
    ("3. Operacion", "schedule_grid_person(+territory)", "Personas + zonas", "#9 #10"),
    ("3. Operacion", "schedule_cell(+territory/break/history)", "Celda dia", "reasignar + auditoria"),
    ("4. Otros", "shift_swap_request", "Intercambio", "#16 sin tope"),
    ("4. Otros", "schedule_novelty", "Novedades", "#18 ligadas a malla"),
    ("4. Otros", "excel_import_*", "Import Excel", "#20 se mantiene"),
]:
    ws2.append(list(r))
for col, w in zip("ABCD", [16, 40, 36, 45]):
    ws2.column_dimensions[col].width = w

ws3 = wb.create_sheet("IDs_externos_GRH")
ws3.append(["Campo_en_malla", "Servicio_GRH", "Tipo_relacion", "Notas"])
for cell in ws3[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E79")
for r in [
    ("company_id", "ms_company_admin / JWT", "UUID logico", "Tenant en todas las tablas"),
    ("employee_id", "ms_employee", "UUID logico", "Personas de malla/celdas"),
    ("area_id", "ms_parametrization", "UUID logico", "front_area_link"),
    ("user_id / *_by", "ms_auth", "UUID logico", "Auditoria y alcance"),
    ("role_id", "ms_auth", "UUID logico", "front_role_scope + RBAC menu"),
    ("Festivos", "ms_parametrization", "API lectura", "No tabla local de cortes (#6)"),
    ("Notificaciones", "ms_notification", "Event types", "Publicacion y cambios (#13)"),
    ("Modulo/permisos", "company-admin + auth", "Seed config", "HU01 HU02 fuera de esta BD"),
]:
    ws3.append(list(r))
for col, w in zip("ABCD", [28, 28, 16, 48]):
    ws3.column_dimensions[col].width = w

xlsx_path = OUT / "PROPUESTA-BD-MALLA-TURNOS.xlsx"
wb.save(xlsx_path)

# DBML
dbml = r'''// Propuesta BD — Malla de Turnos (ms_malla_turnos)
// Post sesion stakeholder 11-sep-2026
// IDs externos GRH = UUID logicos SIN FK cross-database
// Pegar en https://dbdiagram.io

Project malla_turnos {
  database_type: 'PostgreSQL'
  Note: 'BD nueva ms_malla_turnos. Tenant=company_id. Integracion GRH por UUID logicos.'
}

Table operational_front {
  id uuid [pk]
  company_id uuid [not null, note: 'Tenant GRH (logico)']
  code varchar(20) [not null]
  name varchar(120) [not null]
  description text
  builder_profile_label varchar(120)
  publisher_profile_label varchar(120)
  separate_builder_publisher boolean [not null, default: false]
  build_mode varchar(20) [not null, note: 'manual|asistido|automatico']
  period_type varchar(10) [not null, note: 'semana|quincena|mes']
  published_editable boolean [not null, default: false]
  swap_enabled boolean [not null, default: false, note: 'Sin tope mensual solicitudes #16']
  active boolean [not null, default: true]
  created_at timestamptz [not null]
  updated_at timestamptz [not null]
  created_by uuid [note: 'ms_auth user_id logico']
  updated_by uuid
  Indexes {
    (company_id, code) [unique]
  }
}

Table front_capability {
  id uuid [pk]
  company_id uuid [not null]
  front_id uuid [not null, ref: > operational_front.id]
  capability_code varchar(40) [not null, note: 'usa_territorio|usa_campanas|usa_modalidad|usa_sitio|usa_multi_break|usa_multi_zona']
  enabled boolean [not null, default: false]
  config_json jsonb
  Indexes {
    (company_id, front_id, capability_code) [unique]
  }
}

Table front_area_link {
  id uuid [pk]
  company_id uuid [not null]
  front_id uuid [not null, ref: > operational_front.id]
  area_id uuid [not null, note: 'ms_parametrization area_id logico']
  Indexes {
    (company_id, front_id, area_id) [unique]
  }
}

Table front_user_scope {
  id uuid [pk]
  company_id uuid [not null]
  front_id uuid [not null, ref: > operational_front.id]
  user_id uuid [not null, note: 'ms_auth user_id logico']
  Indexes {
    (company_id, front_id, user_id) [unique]
  }
}

Table front_role_scope {
  id uuid [pk]
  company_id uuid [not null]
  front_id uuid [not null, ref: > operational_front.id]
  role_id uuid [not null, note: 'ms_auth role_id logico']
  Indexes {
    (company_id, front_id, role_id) [unique]
  }
}

Table catalog_code_sequence {
  id uuid [pk]
  company_id uuid [not null]
  catalog_type varchar(40) [not null]
  prefix varchar(10)
  next_value int [not null, default: 1]
  Indexes {
    (company_id, catalog_type) [unique]
  }
  Note: 'Codigos auto-consecutivos #3. Sitios pueden ser manual.'
}

Table shift_template {
  id uuid [pk]
  company_id uuid [not null]
  code varchar(20) [not null]
  name varchar(120) [not null]
  start_time time [not null]
  end_time time [not null]
  crosses_midnight boolean [not null, default: false]
  color_hex varchar(7)
  net_hours numeric(5,2)
  multi_break_enabled boolean [not null, default: false, note: '#11']
  active boolean [not null, default: true]
  created_at timestamptz [not null]
  updated_at timestamptz [not null]
  Indexes {
    (company_id, code) [unique]
  }
  Note: 'Catalogo unico. Un turno = un horario; la zona no cambia el turno.'
}

Table shift_template_front {
  id uuid [pk]
  company_id uuid [not null]
  shift_template_id uuid [not null, ref: > shift_template.id]
  front_id uuid [not null, ref: > operational_front.id]
  Indexes {
    (company_id, shift_template_id, front_id) [unique]
  }
}

Table shift_template_day_schedule {
  id uuid [pk]
  company_id uuid [not null]
  shift_template_id uuid [not null, ref: > shift_template.id]
  dow smallint [not null]
  start_time time [not null]
  end_time time [not null]
  crosses_midnight boolean [not null, default: false]
  Indexes {
    (shift_template_id, dow) [unique]
  }
}

Table shift_template_break {
  id uuid [pk]
  company_id uuid [not null]
  shift_template_id uuid [not null, ref: > shift_template.id]
  code varchar(20)
  label varchar(80) [not null]
  duration_minutes int [not null]
  paid boolean [not null, default: false]
  sort_order smallint [not null, default: 0]
  active boolean [not null, default: true]
}

Table cell_state {
  id uuid [pk]
  company_id uuid [not null]
  code varchar(20) [not null]
  name varchar(120) [not null]
  color_hex varchar(7)
  assignable boolean [not null, default: true]
  counts_hours boolean [not null, default: false]
  deducts_productive boolean [not null, default: false]
  requires_support boolean [not null, default: false]
  active boolean [not null, default: true]
  Indexes {
    (company_id, code) [unique]
  }
}

Table work_modality {
  id uuid [pk]
  company_id uuid [not null]
  code varchar(20) [not null]
  name varchar(120) [not null]
  icon varchar(40)
  active boolean [not null, default: true]
  Indexes {
    (company_id, code) [unique]
  }
  Note: '#2 Catalogo GLOBAL. No hay tabla modality_front.'
}

Table attendance_site {
  id uuid [pk]
  company_id uuid [not null]
  code varchar(20) [not null, note: 'Puede ser manual codigo negocio #3']
  name varchar(120) [not null]
  site_type varchar(40)
  address varchar(255)
  active boolean [not null, default: true]
  Indexes {
    (company_id, code) [unique]
  }
}

Table attendance_site_front {
  id uuid [pk]
  company_id uuid [not null]
  site_id uuid [not null, ref: > attendance_site.id]
  front_id uuid [not null, ref: > operational_front.id]
  Indexes {
    (company_id, site_id, front_id) [unique]
  }
}

Table campaign {
  id uuid [pk]
  company_id uuid [not null]
  code varchar(20) [not null]
  name varchar(120) [not null]
  front_id uuid [ref: > operational_front.id]
  valid_from date
  valid_to date
  active boolean [not null, default: true]
  Indexes {
    (company_id, code) [unique]
  }
  Note: 'No se asigna en paso 1 de crear malla #17; va en celda.'
}

Table territory_level {
  id uuid [pk]
  company_id uuid [not null]
  level_order smallint [not null]
  name varchar(80) [not null]
  Indexes {
    (company_id, level_order) [unique]
  }
}

Table territory_node {
  id uuid [pk]
  company_id uuid [not null]
  level_id uuid [not null, ref: > territory_level.id]
  parent_id uuid [ref: > territory_node.id]
  code varchar(20) [not null]
  name varchar(120) [not null]
  active boolean [not null, default: true]
  Indexes {
    (company_id, code) [unique]
  }
}

Table validation_rule {
  id uuid [pk]
  company_id uuid [not null]
  code varchar(40) [not null]
  scope varchar(20) [not null, note: 'empresa|frente']
  front_id uuid [ref: > operational_front.id]
  param_value varchar(40)
  param_unit varchar(40)
  severity varchar(20) [not null]
  priority int [not null, default: 100]
  active boolean [not null, default: true]
  notes text
  Note: 'Reglas internas sin pantallas de cobertura/compensatorio #4 #5'
}

Table hour_classification_band {
  id uuid [pk]
  company_id uuid [not null]
  code varchar(20) [not null]
  name varchar(120) [not null]
  band_type varchar(40) [not null]
  start_time time
  end_time time
  crosses_midnight boolean [not null, default: false]
  applies_sunday boolean [not null, default: false]
  applies_holiday boolean [not null, default: false]
  sort_order smallint [not null, default: 0]
  active boolean [not null, default: true]
  Indexes {
    (company_id, code) [unique]
  }
  Note: 'Franjas recargo #8. Festivos por API parametrization.'
}

Table rotation_pattern {
  id uuid [pk]
  company_id uuid [not null]
  front_id uuid [not null, ref: > operational_front.id]
  code varchar(20) [not null]
  name varchar(120) [not null]
  cycle_length smallint [not null]
  auto_apply_on_period_close boolean [not null, default: true, note: '#12']
  active boolean [not null, default: true]
}

Table rotation_pattern_step {
  id uuid [pk]
  company_id uuid [not null]
  pattern_id uuid [not null, ref: > rotation_pattern.id]
  step_order smallint [not null]
  shift_template_id uuid [ref: > shift_template.id]
  cell_state_id uuid [ref: > cell_state.id]
}

Table schedule_grid {
  id uuid [pk]
  company_id uuid [not null]
  front_id uuid [not null, ref: > operational_front.id]
  name varchar(160)
  period_start date [not null]
  period_end date [not null]
  period_type varchar(10) [not null]
  status varchar(20) [not null, note: 'borrador|revision|publicada|cerrada|rechazada']
  build_mode varchar(20) [not null]
  owner_user_id uuid [note: 'ms_auth']
  publisher_user_id uuid [note: 'ms_auth']
  rejection_reason text
  published_at timestamptz
  closed_at timestamptz
  version int [not null, default: 1]
  created_at timestamptz [not null]
  updated_at timestamptz [not null]
  Indexes {
    (company_id, front_id, period_start) [name: 'ix_grid_list']
  }
  Note: 'Sin campaign_id en cabecera #17'
}

Table schedule_grid_person {
  id uuid [pk]
  company_id uuid [not null]
  grid_id uuid [not null, ref: > schedule_grid.id]
  employee_id uuid [not null, note: 'ms_employee logico']
  sort_order int [not null, default: 0]
  modality_id uuid [ref: > work_modality.id]
  site_id uuid [ref: > attendance_site.id]
  Indexes {
    (grid_id, employee_id) [unique]
  }
}

Table schedule_grid_person_territory {
  id uuid [pk]
  company_id uuid [not null]
  grid_person_id uuid [not null, ref: > schedule_grid_person.id]
  territory_node_id uuid [not null, ref: > territory_node.id]
  is_primary boolean [not null, default: false]
  Indexes {
    (grid_person_id, territory_node_id) [unique]
  }
  Note: 'Multi zona/SPT por persona #10 #14'
}

Table schedule_cell {
  id uuid [pk]
  company_id uuid [not null]
  grid_id uuid [not null, ref: > schedule_grid.id]
  employee_id uuid [not null, note: 'ms_employee logico']
  work_date date [not null]
  slot smallint [not null, default: 0]
  shift_template_id uuid [ref: > shift_template.id]
  cell_state_id uuid [ref: > cell_state.id]
  modality_id uuid [ref: > work_modality.id]
  site_id uuid [ref: > attendance_site.id]
  campaign_id uuid [ref: > campaign.id]
  note text
  locked_attrs jsonb
  origin varchar(40)
  updated_at timestamptz [not null]
  updated_by uuid
  Indexes {
    (grid_id, employee_id, work_date, slot) [unique]
    (company_id, employee_id, work_date) [name: 'ix_cell_who']
  }
}

Table schedule_cell_territory {
  id uuid [pk]
  company_id uuid [not null]
  cell_id uuid [not null, ref: > schedule_cell.id]
  territory_node_id uuid [not null, ref: > territory_node.id]
  Indexes {
    (cell_id, territory_node_id) [unique]
  }
}

Table schedule_cell_break {
  id uuid [pk]
  company_id uuid [not null]
  cell_id uuid [not null, ref: > schedule_cell.id]
  template_break_id uuid [ref: > shift_template_break.id]
  label varchar(80) [not null]
  duration_minutes int [not null]
  paid boolean [not null, default: false]
  sort_order smallint [not null, default: 0]
  Note: 'Breaks repetibles / especiales #11'
}

Table schedule_cell_history {
  id uuid [pk]
  company_id uuid [not null]
  cell_id uuid [not null, ref: > schedule_cell.id]
  grid_id uuid [not null]
  employee_id uuid [not null]
  work_date date [not null]
  action varchar(40) [not null, note: 'reasignar|publicar|intercambio|novedad|import']
  before_json jsonb
  after_json jsonb
  reason text
  origin varchar(40)
  changed_by uuid [not null, note: 'ms_auth']
  changed_at timestamptz [not null]
  Indexes {
    (cell_id, changed_at) [name: 'ix_cell_hist']
  }
}

Table shift_swap_request {
  id uuid [pk]
  company_id uuid [not null]
  front_id uuid [not null, ref: > operational_front.id]
  grid_id uuid [not null, ref: > schedule_grid.id]
  requester_employee_id uuid [not null]
  counterpart_employee_id uuid [not null]
  requester_cell_id uuid [not null, ref: > schedule_cell.id]
  counterpart_cell_id uuid [not null, ref: > schedule_cell.id]
  status varchar(30) [not null]
  rule_warnings jsonb
  decision_reason text
  decided_by uuid
  decided_at timestamptz
  created_at timestamptz [not null]
  updated_at timestamptz [not null]
  Note: 'Sin max_requests_per_month #16'
}

Table schedule_novelty {
  id uuid [pk]
  company_id uuid [not null]
  grid_id uuid [not null, ref: > schedule_grid.id]
  cell_id uuid [ref: > schedule_cell.id]
  employee_id uuid [not null]
  work_date date [not null]
  cell_state_id uuid [ref: > cell_state.id]
  description text
  source varchar(40) [not null]
  created_by uuid
  created_at timestamptz [not null]
  Note: 'Novedades vinculadas a malla #18'
}

Table excel_import_batch {
  id uuid [pk]
  company_id uuid [not null]
  front_id uuid [ref: > operational_front.id]
  import_type varchar(40) [not null]
  file_name varchar(255)
  status varchar(20) [not null]
  summary_json jsonb
  uploaded_by uuid
  created_at timestamptz [not null]
  Note: 'Plantilla Excel se mantiene #20'
}

Table excel_import_row {
  id uuid [pk]
  company_id uuid [not null]
  batch_id uuid [not null, ref: > excel_import_batch.id]
  row_number int [not null]
  payload_json jsonb [not null]
  action varchar(40)
  status varchar(20) [not null]
  message text
}
'''

dbml_path = OUT / "PROPUESTA-BD-MALLA-TURNOS.dbml"
dbml_path.write_text(dbml, encoding="utf-8")

print("OK CSV:", csv_path)
print("OK XLSX:", xlsx_path)
print("OK DBML:", dbml_path)
print("Filas columnas:", len(rows))
