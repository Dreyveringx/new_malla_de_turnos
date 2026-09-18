# -*- coding: utf-8 -*-
"""Genera PROPUESTA-BD-MALLA-TURNOS-GUIA-TABLAS.docx (v2 — post revision transversal)."""
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    import subprocess
    import sys

    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

OUT = Path(__file__).resolve().parent / "PROPUESTA-BD-MALLA-TURNOS-GUIA-TABLAS.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def set_cell_shading(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
    return h


def p(text, bold=False, italic=False, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    return para


def bullet(text, bold_prefix=None):
    para = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = para.add_run(bold_prefix)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r2 = para.add_run(text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(11)
    else:
        run = para.add_run(text)
        run.font.size = Pt(11)
        run.font.name = "Calibri"
    return para


def table_header(cells, hex_color="1A3A5C"):
    for c in cells:
        set_cell_shading(c, hex_color)
        for para in c.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True


def add_meta(base, decision, color_rgb):
    meta = doc.add_paragraph()
    r1 = meta.add_run("Base de datos: ")
    r1.bold = True
    meta.add_run(f"{base}    ")
    r2 = meta.add_run("Decision: ")
    r2.bold = True
    r3 = meta.add_run(decision)
    r3.bold = True
    r3.font.color.rgb = color_rgb


# ---- Portada ----
t = doc.add_heading("Guia de tablas — Propuesta BD Malla de Turnos (v2)", 0)
for run in t.runs:
    run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)

p(
    "Actualizado con la revision de transversalidad, reuso GRH y eliminacion de reglas "
    "(cobertura, compensatorio, restricciones, normas). Separacion clara por base de datos.",
    italic=True,
)
p("Companion: PROPUESTA-BD-MALLA-TURNOS-v2.xlsx | .dbml | DECISIONES.md")
p("")

add_heading("1. Idea clave", 1)
p(
    "Los catalogos reutilizables viven en ms_parametrization. El dominio de malla "
    "(frente, grilla, celdas, swap, novedades, import) vive en ms_malla_turnos. "
    "Historial de cambios va a ms_audit. Entre BDs solo UUID logicos, sin FK cross-database."
)
p(
    "Roles y permisos = RBAC de ms_auth (sin front_role_scope). "
    "Estados de celda = entity_status existente. "
    "work_schedule es jornada contractual: NO se reutiliza como plantilla de turno de malla."
)

add_heading("2. Resumen por base de datos", 1)
tbl = doc.add_table(rows=1, cols=3)
tbl.style = "Table Grid"
hdr = tbl.rows[0].cells
hdr[0].text = "Base de datos"
hdr[1].text = "Que va ahi"
hdr[2].text = "Ejemplos"
table_header(hdr)
for a, b, c in [
    (
        "ms_parametrization",
        "Catalogos transversales NUEVOS + reuso entity_status",
        "shift_template, break_type, work_modality, attendance_site, operational_label, territory, hour_classification_band, catalog_code_sequence",
    ),
    (
        "ms_malla_turnos",
        "Dominio operativo del modulo",
        "operational_front, schedule_grid/cell, rotacion, swap, novedades, import Excel, vinculos frente",
    ),
    ("ms_audit", "Historial / auditoria", "Reemplaza schedule_cell_history; UI consulta audit/timeline"),
    (
        "Otros GRH (solo UUID)",
        "No se duplican",
        "company, employee, areas, positions, auth, festivos, notificaciones",
    ),
]:
    row = tbl.add_row().cells
    row[0].text = a
    row[1].text = b
    row[2].text = c

p("")

add_heading("3. Que ya existe en GRH y como se reutiliza", 1)
reuse = doc.add_table(rows=1, cols=3)
reuse.style = "Table Grid"
rh = reuse.rows[0].cells
rh[0].text = "Concepto"
rh[1].text = "Donde esta"
rh[2].text = "Como lo usa Malla"
table_header(rh)
for a, b, c in [
    ("Empresa (tenant)", "ms_company_admin / JWT", "company_id UUID logico"),
    ("Empleado", "ms_employee", "employee_id; datos por API"),
    ("Areas / cargos", "organizational_areas, positions", "front_area_link.area_id"),
    ("Usuarios / roles / permisos", "ms_auth RBAC", "Sin front_role_scope"),
    ("Estados de celda", "entity_status", "entity_type = malla_cell_state"),
    ("Festivos / calendario", "company_calendar, national_holidays", "Lectura API para clasificacion de horas"),
    ("Jornada contractual", "work_schedule + day + shift + schedule_assignment", "NO usar como plantilla de malla"),
    ("Modalidad de estudio", "study_modalities", "NO sirve para modalidad laboral"),
    ("Tipos de zona", "zone_types", "NO es arbol de cobertura operativa"),
    ("Auditoria", "ms_audit", "Historial de celdas/cambios"),
    ("Notificaciones", "ms_notification", "Publicacion y cambios"),
]:
    row = reuse.add_row().cells
    row[0].text = a
    row[1].text = b
    row[2].text = c

p("")

# ---- Param ----
add_heading("4. Catalogos en ms_parametrization", 1)
p("Nuevos (transversales para Malla y otros modulos) o reuso explicito.")

param_tables = [
    (
        "catalog_code_sequence",
        "NUEVA transversal",
        "Secuencia de codigos auto-consecutivos por (company_id, catalog_type). "
        "Sirve para malla y para cualquier catalogo GRH que hoy pide codigo manual.",
        "API/servicio en parametrization. Malla (y otros) piden el siguiente codigo al crear registros.",
    ),
    (
        "shift_template",
        "NUEVA transversal (NO es work_schedule)",
        "Plantilla de turno operativo (Manana A, Noche B): codigo, horario, color, multi-break. "
        "work_schedule es jornada contractual semanal (horas base, CPM). Dominios distintos → se deja esta tabla.",
        "Malla guarda shift_template_id (UUID logico) en celdas y en shift_template_front.",
    ),
    (
        "break_type",
        "NUEVA transversal",
        "Tipos de pausa (almuerzo, break, lactancia...) reutilizables por cualquier modulo.",
        "Referenciado por shift_template_break y por schedule_cell_break.",
    ),
    (
        "shift_template_break",
        "NUEVA transversal",
        "Breaks definidos en una plantilla de turno (1..N), con tipo y minutos.",
        "Al asignar un turno a una celda se pueden copiar/ajustar en schedule_cell_break.",
    ),
    (
        "work_modality",
        "NUEVA transversal",
        "Modalidad laboral global (presencial, remoto, hibrido). No confundir con study_modalities.",
        "UUID en schedule_grid_person / schedule_cell. Sin modality_front (global).",
    ),
    (
        "attendance_site",
        "NUEVA transversal",
        "Sitios/sedes de asistencia. No existia en param ni company-admin. Sin attendance_site_front: "
        "si el frente tiene usa_sitio, usa el catalogo de la empresa.",
        "UUID site_id en persona/celda de malla.",
    ),
    (
        "operational_label",
        "NUEVA transversal (antes campaign)",
        "Etiquetas informativas (campana, cliente, tag). Catalogo transversal.",
        "En celda: label_id. No va en cabecera de malla (#17).",
    ),
    (
        "territory",
        "NUEVA transversal (arbol)",
        "Una sola tabla con parent_id (hijos anidados). Reemplaza territory_level + territory_node. "
        "level_label texto libre (Zona, SPT). No usar zone_types ni organizational_areas.",
        "UUID territory_id en schedule_grid_person_territory / schedule_cell_territory.",
    ),
    (
        "hour_classification_band",
        "NUEVA transversal",
        "Franjas para clasificar horas (ordinaria, nocturna, dominical, festiva, extras). "
        "Si hay varios turnos el mismo dia (slots), cada celda se clasifica. Festivos ya existen. "
        "No calcula nomina ni cortes (#6).",
        "Reporteria/export desde malla u otros modulos.",
    ),
    (
        "entity_status (REUSO)",
        "YA EXISTE — no crear cell_state",
        "Usar entity_type = 'malla_cell_state' para Trabajo, Descanso, VAC, INC, etc. "
        "Ya tiene name, color, sort_order, is_active, is_final_status.",
        "schedule_cell.cell_state_id = entity_status.id (UUID logico).",
    ),
]

for name, decision, purpose, how in param_tables:
    add_heading(name, 2)
    add_meta("ms_parametrization", decision, RGBColor(0x2E, 0x7D, 0x32))
    p("Para que sirve", bold=True)
    p(purpose)
    p("Como se reutiliza / consume", bold=True)
    p(how)

# ---- Malla ----
add_heading("5. Dominio en ms_malla_turnos", 1)

malla_tables = [
    (
        "operational_front",
        "Frente operativo: periodo (semana/quincena/mes), edicion publicada, swap, modo de armado.",
    ),
    (
        "front_capability",
        "Flags por frente: usa_territorio, usa_labels, usa_modalidad, usa_sitio, usa_multi_break, usa_multi_zona.",
    ),
    (
        "front_area_link",
        "Vinculo frente ↔ organizational_areas (UUID logico).",
    ),
    (
        "front_user_scope",
        "Alcance por usuario (que frentes ve). Roles → RBAC; no hay front_role_scope.",
    ),
    (
        "shift_template_front",
        "Que plantillas de turno (param) estan habilitadas en el frente.",
    ),
    (
        "rotation_pattern / rotation_pattern_step",
        "Patrones de rotacion al cerrar periodo (#12). Referencian shift_template / entity_status por UUID.",
    ),
    (
        "schedule_grid",
        "Cabecera de malla (frente + periodo + estado). Sin etiqueta en cabecera.",
    ),
    (
        "schedule_grid_person",
        "Personas de la malla (employee_id). Modalidad/sitio opcionales del periodo.",
    ),
    (
        "schedule_grid_person_territory",
        "Multi-zona/SPT de la persona en el periodo (#10/#14).",
    ),
    (
        "schedule_cell",
        "Celda dia x persona (slot). Referencias logicas a turnos, estados, modalidad, sitio, label.",
    ),
    (
        "schedule_cell_territory",
        "Override de territorio en la celda.",
    ),
    (
        "schedule_cell_break",
        "Breaks reales de la celda (especiales/repetidos #11).",
    ),
    (
        "shift_swap_request",
        "Intercambio sin tope mensual. Sin motor de cobertura/compensatorio.",
    ),
    (
        "schedule_novelty",
        "Novedades vinculadas a la malla (#18).",
    ),
    (
        "excel_import_batch / excel_import_row",
        "Staging de importacion Excel (#20).",
    ),
]

for name, purpose in malla_tables:
    add_heading(name, 2)
    add_meta("ms_malla_turnos", "DOMINIO — se crea aqui", RGBColor(0x1F, 0x4E, 0x79))
    p("Para que sirve", bold=True)
    p(purpose)

# ---- Audit ----
add_heading("6. Auditoria (ms_audit)", 1)
add_heading("audit_log / timeline", 2)
add_meta("ms_audit", "REUSO — no crear schedule_cell_history", RGBColor(0x6A, 0x1B, 0x9A))
p(
    "Cada cambio de celda, publicacion o swap emite evento a ms_audit (y/o timeline). "
    "La pantalla de historial de la malla consulta ese servicio. No hay tabla local de historial."
)

# ---- Eliminados ----
add_heading("7. Eliminado del diseno", 1)
drop = doc.add_table(rows=1, cols=2)
drop.style = "Table Grid"
dh = drop.rows[0].cells
dh[0].text = "Tabla / concepto"
dh[1].text = "Motivo"
table_header(dh, "B71C1C")
for a, b in [
    ("front_role_scope", "Roles con RBAC ms_auth"),
    ("attendance_site_front", "Basta capability usa_sitio + catalogo empresa"),
    ("shift_template_day_schedule", "Un turno = un horario; variacion diaria es work_schedule contractual"),
    ("territory_level + territory_node", "Reemplazados por territory (parent_id)"),
    ("campaign", "Renombrado a operational_label"),
    ("cell_state (tabla propia)", "Reutilizar entity_status"),
    ("validation_rule", "Sin reglas operativas / cobertura / compensatorio / restricciones / normas en BD"),
    ("coverage / compensatory / assignment_restriction / operational_norm", "Fuera de alcance"),
    ("payroll_cutoff", "Sesion #6"),
    ("work_modality_front", "Modalidad global"),
    ("swap_monthly_limit", "Sesion #16"),
    ("schedule_cell_history", "Pasa a ms_audit"),
    ("catalog_code_sequence en malla", "Movida a parametrization (transversal)"),
]:
    row = drop.add_row().cells
    row[0].text = a
    row[1].text = b

p("")

add_heading("8. MVP sugerido", 1)
p("Parametrization (minimo)", bold=True)
for x in [
    "catalog_code_sequence",
    "shift_template (+ break_type / shift_template_break si hay breaks)",
    "entity_status con entity_type=malla_cell_state",
    "work_modality (si el frente usa modalidad)",
]:
    bullet(x)

p("Malla (minimo)", bold=True)
for x in [
    "operational_front + front_capability + front_area_link",
    "shift_template_front",
    "schedule_grid + schedule_grid_person + schedule_cell",
    "front_user_scope (si hace falta alcance por usuario)",
]:
    bullet(x)

p("Diferir", bold=True)
for x in [
    "territory* / labels / attendance_site si el frente no los usa",
    "rotation_pattern*, swap, novelty, excel_import_*",
    "hour_classification_band hasta reporteria de recargos",
]:
    bullet(x)

add_heading("9. Archivos companion", 1)
bullet("PROPUESTA-BD-MALLA-TURNOS-v2.xlsx — columnas por BD + hoja Decisiones", bold_prefix="Excel: ")
bullet("PROPUESTA-BD-MALLA-TURNOS.dbml — pegar en dbdiagram.io", bold_prefix="Diagrama: ")
bullet("PROPUESTA-BD-MALLA-TURNOS-DECISIONES.md — resumen corto", bold_prefix="Markdown: ")
bullet("este Word — explicacion tabla a tabla v2", bold_prefix="Guia: ")

# Save with fallback if locked
try:
    doc.save(OUT)
    print("OK", OUT)
except PermissionError:
    alt = OUT.with_name("PROPUESTA-BD-MALLA-TURNOS-GUIA-TABLAS-v2.docx")
    doc.save(alt)
    print("OK (fallback)", alt)
print("bytes", OUT.stat().st_size if OUT.exists() else "n/a")
